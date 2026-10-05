"""Readable Excel presentation; original field names remain the CSV contract."""
from datetime import date, datetime
from decimal import Decimal
from math import ceil
import re
from zoneinfo import ZoneInfo

from openpyxl.styles import Alignment, Font, PatternFill

SHEETS = {
    'Summary': ('О файле', 'Магазин, период и время выгрузки.'),
    'Daily': ('Показатели по дням', 'Сохранённые ежедневные показатели WB и Ozon.'),
    'Products': ('Товары', 'Названия, артикулы и заказы за период.'),
    'Inventory': ('Остатки', 'Последние снимки количества на складах.'),
    'Inbound': ('Поставки в пути', 'Товары, статусы и даты прибытия.'),
    'Promotions': ('Акции', 'Календарь акций и товары участников.'),
    'Supply': ('План поставок', 'Рекомендации по пополнению.'),
    'ForecastQuality': ('Точность прогноза', 'Прогноз и реальные заказы.'),
    'Calibration': ('Поправки к запасу', 'Дополнительный запас на риск.'),
    'Actions': ('Задачи на сегодня', 'Текущие рекомендации.'),
    'ActionHistory': ('История задач', 'Предыдущие подсказки и ваши отметки.'),
    'Finance': ('Финансовая сводка', 'Продажи, начисления и удержания по источникам.'),
    'Advertising': ('Реклама', 'Рекламные расходы и продажи.'),
    'Management': ('Результат магазина', 'Предварительная оценка по известным расходам.'),
    'Reconciliation': ('Сверка данных', 'Заказы, доставка и начисления.'),
    'Costs': ('Себестоимость', 'Стоимость одной штуки и дата её действия.'),
    'AccrualSummary': ('Итог', 'Итоги и полнота выбранных начислений.'),
    'OzonOperations': ('Операции Ozon', 'Сохранённые финансовые операции.'),
    'WbReports': ('Отчёты WB', 'Итоги отчётов WB с настоящими периодами.'),
}

TITLES = {
    'key':'Показатель', 'value':'Значение', 'marketplace':'Площадка',
    'metric_key':'Код показателя', 'data_date':'Дата', 'date':'Дата сохранения',
    'name':'Название', 'product_name':'Название товара', 'marketplace_sku':'Артикул площадки',
    'sku':'Артикулы', 'internal_sku':'Общее имя товара', 'offer_id':'Артикул продавца',
    'product_id':'Номер товара в боте', 'listing_id':'Номер карточки в боте', 'connection_id':'Номер подключения',
    'unit':'Единица', 'is_preliminary':'Предварительные данные', 'as_of':'Данные на',
    'fetched_at':'Дата получения', 'captured_at':'Снимок остатков', 'updated_at':'Обновлено',
    'ordered_units':'Заказано, шт.', 'ordered_revenue':'Сумма заказов, ₽',
    'available_units':'Доступно, шт.', 'reserved_units':'Зарезервировано, шт.',
    'fulfillment_scheme':'Схема склада', 'cost_price':'Себестоимость, ₽',
    'effective_date':'Дата начала действия', 'source':'Источник',
    'financial_sales':'Продажи по финансовому источнику, ₽',
    'commission':'Комиссия, ₽', 'logistics':'Логистика, ₽', 'storage':'Хранение, ₽',
    'acceptance':'Приёмка, ₽', 'services':'Услуги и прочие удержания, ₽',
    'penalties':'Штрафы, ₽', 'compensation':'Компенсации и доплаты, ₽',
    'marketplace_net':'Начисления после удержаний, ₽', 'goods_payable':'К перечислению за товар, ₽',
    'bank_payment':'К оплате по отчёту, ₽', 'other_components':'Другие составляющие итога, ₽',
    'ads_already_in_services':'Реклама в услугах, ₽', 'finance_ad_spend':'Реклама в начислениях, ₽',
    'source_run_id':'Номер сохранённого ответа', 'operation_index':'Строка операции',
    'operation_id':'Номер операции', 'unit_number':'Номер расчётной единицы',
    'operation_date':'Дата операции', 'category':'Категория', 'posting_number':'Номер отправления',
    'fee_type_ids':'Коды удержаний', 'verified_at':'Источник проверен',
    'report_id':'Номер отчёта WB', 'report_type':'Тип отчёта WB',
    'report_from':'Период отчёта с', 'report_to':'Период отчёта по',
    'spend':'Расходы на рекламу, ₽', 'attributed_sales':'Продажи по рекламе, ₽',
    'ad_spend':'Расходы на рекламу, ₽', 'orders':'Заказы, шт.', 'clicks':'Клики',
    'impressions':'Показы', 'drr_pct':'ДРР, %', 'roas':'Отдача рекламы', 'level':'Уровень данных',
    'estimated_cogs':'Учтённая себестоимость, ₽', 'estimated_order_cogs':'Себестоимость заказов, ₽',
    'cogs_coverage_pct':'Покрытие себестоимостью, %', 'estimated_result':'Предварительный результат, ₽',
    'marketplace_expenses':'Расходы площадки, ₽', 'external_supply_id':'Номер поставки',
    'status':'Статус', 'planned_at':'Плановая дата', 'arrival_at':'Прибытие',
    'warehouse_name':'Склад', 'planned_units':'В плане, шт.', 'accepted_units':'Принято, шт.',
    'remaining_units':'Осталось принять, шт.', 'external_promotion_id':'Номер акции',
    'promotion_name':'Акция', 'promo_type':'Тип акции', 'start_at':'Начало', 'end_at':'Окончание',
    'in_action':'Участвует в акции', 'base_price':'Базовая цена, ₽', 'promo_price':'Цена акции, ₽',
    'discount_pct':'Скидка, %', 'abc_class':'Класс ABC', 'xyz_class':'Класс XYZ',
    'avg_daily_units':'Средние заказы в день, шт.', 'forecast_daily_units':'Прогноз в день, шт.',
    'forecast_next_7_units':'Прогноз на 7 дней, шт.', 'seasonality_strength':'Сила сезонности',
    'bias_correction':'Поправка прогноза', 'promo_factor':'Влияние акций', 'promo_days':'Дни акций',
    'trend_pct':'Изменение заказов, %', 'inbound_units':'В пути, шт.', 'effective_units':'Запас с поставками, шт.',
    'days_cover':'Дни запаса', 'lead_time_days':'Доставка, дн.', 'lead_buffer_days':'Поправка доставки, дн.',
    'effective_lead_time_days':'Доставка с поправкой, дн.', 'safety_stock_days':'Страховой запас, дн.',
    'safety_buffer_days':'Поправка страхового запаса, дн.', 'effective_safety_stock_days':'Страховой запас с поправкой, дн.',
    'target_stock_days':'Целевой запас, дн.', 'calibration_confidence':'Надёжность поправок',
    'reorder_point_units':'Порог пополнения, шт.', 'target_units':'Цель, шт.',
    'recommended_order_units':'Рекомендовано заказать, шт.', 'pack_size':'Упаковка, шт.',
    'min_order_qty':'Минимальная партия, шт.', 'confidence':'Надёжность', 'history_days':'Дни истории',
    'inventory_as_of':'Дата остатков', 'inventory_age_days':'Возраст остатков, дн.', 'inventory_stale':'Остатки устарели',
    'samples':'Число наблюдений', 'predicted_units':'Прогноз, шт.', 'actual_units':'Фактически, шт.',
    'mae_units':'Средняя ошибка, шт.', 'wape_pct':'Ошибка прогноза, %', 'bias_pct':'Смещение прогноза, %',
    'forecast_samples':'Наблюдения прогноза', 'forecast_wape_pct':'Ошибка прогноза, %',
    'forecast_bias_pct':'Смещение прогноза, %', 'inventory_samples':'Наблюдения остатков',
    'zero_stock_rate_pct':'Доля нулевых остатков, %', 'inbound_delay_samples':'Наблюдения поставок',
    'avg_inbound_delay_days':'Средняя задержка, дн.', 'p75_inbound_delay_days':'Задержка в 75% случаев, дн.',
    'reasons':'Причины', 'action_key':'Код задачи', 'priority':'Приоритет', 'title':'Задача',
    'detail':'Подробности', 'evidence':'Исходные данные', 'hint':'Что сделать',
    'event_type':'Действие', 'happened_at':'Дата действия', 'actor_user_id':'Telegram ID сотрудника',
    'note':'Пояснение', 'shop_id':'Номер магазина', 'shop_name':'Магазин',
    'period_start':'Период с', 'period_end':'Период по', 'generated_at_utc':'Создано',
    'credential_profile':'Профиль подключения', 'schema_version':'Версия схемы',
    'loaded_days':'Загружено дней', 'expected_days':'Дней в периоде', 'operations':'Операций',
    'reports':'Финансовых отчётов', 'warning':'Предупреждение', 'loaded_dates':'Загруженные даты',
    'id':'Номер записи', 'as_of_date':'Дата рекомендаций', 'created_at':'Создано',
    'current_status':'Текущий статус', 'snoozed_until':'Отложено до',
    'evidence_json':'Исходные данные задачи', 'cancelled_units':'Отменено, шт.',
    'posting_units':'В отправлениях, шт.', 'sale_units':'Продано, шт.', 'return_units':'Возвраты, шт.',
    'finance_gross':'Финансовые продажи для сверки, ₽', 'ad_attributed_sales':'Продажи по рекламе, ₽',
}
MONEY_FIELDS = {key for key, title in TITLES.items() if title.endswith(', ₽')}
DATE_FIELDS = {'date','data_date','operation_date','effective_date','report_from','report_to',
               'period_start','period_end','as_of_date'}
TIME_FIELDS = {'verified_at','fetched_at','captured_at','updated_at','as_of','planned_at','arrival_at',
               'start_at','end_at','inventory_as_of','happened_at','generated_at_utc','created_at',
               'acknowledged_at','snoozed_until','resolved_at'}
TEXT_FIELDS = {'sku','internal_sku','marketplace_sku','offer_id','source_run_id','operation_id',
               'posting_number','unit_number','report_id','external_supply_id','external_promotion_id',
               'product_id','listing_id','connection_id','actor_user_id','telegram_user_id'}


def title(key):
    return TITLES.get(key, 'Служебное поле: ' + str(key))


def headers(rows):
    return list(dict.fromkeys(key for row in rows for key in row))


def safe_value(value):
    if isinstance(value, str) and value.lstrip(' \t\r\n').startswith(('=','+','-','@')):
        return "'" + value
    return value


def excel_value(key, value, timezone):
    if value is None:return None
    if key in TEXT_FIELDS:return safe_value(str(value))
    if key in DATE_FIELDS | TIME_FIELDS and isinstance(value, str) and value:
        try:
            if len(value) == 10:return date.fromisoformat(value)
            stamp = datetime.fromisoformat(value.replace('Z', '+00:00'))
            if stamp.tzinfo is not None:stamp = stamp.astimezone(ZoneInfo(timezone)).replace(tzinfo=None)
            return stamp
        except ValueError:pass
    if isinstance(value, Decimal):return float(value)
    return safe_value(value)


def write_sheet(workbook, name, rows, *, timezone='Europe/Moscow', keys=None):
    ws = workbook.create_sheet(SHEETS.get(name, (name, ''))[0][:31])
    keys = keys or headers(rows)
    if not keys:
        ws.append(['Нет строк в сохранённых данных'])
        return ws
    ws.append([title(key) for key in keys])
    for row in rows:
        ws.append([title(row[key]) if name == 'Summary' and key == 'key' else
                   excel_value(row['key'] if name == 'Summary' and key == 'value' else key,
                               row.get(key), timezone) for key in keys])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='244761')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 45
    for source, row in zip(rows, ws.iter_rows(min_row=2)):
        for key, cell in zip(keys, row):
            semantic_key = source.get('key') if name == 'Summary' and key == 'value' else key
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=semantic_key not in TEXT_FIELDS or key == 'sku')
            if cell.row % 2 == 0:cell.fill = PatternFill('solid', fgColor='F0F5F8')
            if semantic_key in TEXT_FIELDS:cell.number_format = '@'
            elif isinstance(cell.value, datetime):cell.number_format = 'dd.mm.yyyy hh:mm'
            elif isinstance(cell.value, date):cell.number_format = 'dd.mm.yyyy'
            elif (semantic_key in MONEY_FIELDS or key == 'value' and (
                    source.get('unit') == 'RUB' or source.get('metric_key') in MONEY_FIELDS)) and isinstance(cell.value, (int,float)):
                cell.number_format = '#,##0.00 "₽"'
            elif isinstance(cell.value, (int,float)) and not isinstance(cell.value,bool):
                cell.number_format = '#,##0.00' if key.endswith('_pct') or isinstance(cell.value,float) else '0'
    for col, key in zip(ws.columns, keys):
        sample = [len(str(cell.value or '')) for cell in col[:200]]
        ws.column_dimensions[col[0].column_letter].width = min(42, max(18, len(title(key))+2, *sample))
    for row in ws.iter_rows(min_row=2):
        lines = max((sum(max(1, ceil(len(line) / max(1, ws.column_dimensions[cell.column_letter].width - 2)))
                         for line in str(cell.value or '').split('\n'))
                     if cell.alignment.wrap_text else 1 for cell in row), default=1)
        ws.row_dimensions[row[0].row].height = min(150, max(18, lines * 15))
    return ws


def description_rows(tables):
    rows = []
    for name, data in tables.items():
        for key in headers(data):
            rows.append({'sheet':SHEETS.get(name,(name,''))[0], 'field':key, 'label':title(key),
                'meaning': 'Сумма по отчёту площадки; поступление проверяется в банке.' if key == 'bank_payment' else
                    ('Реклама уже учтена в услугах.' if key == 'ads_already_in_services' else
                     ('Номер сохранён как текст, включая начальные нули.' if key in TEXT_FIELDS else title(key)))})
        if name == 'Summary':
            for item in data:
                key = item['key']
                rows.append({'sheet':'О файле','field':key,'label':title(key),'meaning':title(key)})
    return rows


def add_descriptions(workbook, tables):
    ws = workbook.create_sheet('Описание полей')
    ws.append(['Лист', 'Код поля в CSV', 'Название в Excel', 'Пояснение'])
    for row in description_rows(tables):ws.append([safe_value(row[k]) for k in ('sheet','field','label','meaning')])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='244761')
    for col,width in zip('ABCD',(24,30,42,65)):ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')


def export_filename(prefix, shop_name, start, end, suffix):
    shop = re.sub(r'[^\w\-]', '_', str(shop_name), flags=re.UNICODE).strip('_')[:50] or 'Магазин'
    a=date.fromisoformat(start).strftime('%d-%m-%Y');b=date.fromisoformat(end).strftime('%d-%m-%Y')
    return f'{prefix}_{shop}_{a}' + (f'_{b}' if a != b else '') + suffix
