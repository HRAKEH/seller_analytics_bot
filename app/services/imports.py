"""CSV/XLSX imports for product costs and cross-marketplace product mapping."""
from __future__ import annotations
import csv
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from app.storage import Repository


@dataclass(frozen=True)
class ImportSummary:
    batch_id: int
    filename: str
    total_rows: int
    applied_rows: int
    errors: tuple[dict[str,Any], ...]

    @property
    def status(self) -> str:
        if self.applied_rows == 0 and self.errors: return 'failed'
        if self.errors: return 'partial'
        return 'success'


def _headers(row: dict[str,Any]) -> dict[str,Any]:
    return {str(k or '').strip().lower():v for k,v in row.items()}


def _read_csv(path: Path) -> list[dict[str,Any]]:
    raw=path.read_text(encoding='utf-8-sig')
    sample=raw[:4096]
    try: dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
    except csv.Error: dialect=csv.excel
    return [dict(r) for r in csv.DictReader(raw.splitlines(),dialect=dialect)]


def _read_xlsx(path: Path) -> list[dict[str,Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError('Для XLSX установите openpyxl.') from exc
    wb=load_workbook(path,read_only=True,data_only=True)
    ws=wb.active
    iterator=ws.iter_rows(values_only=True)
    try: header=[str(v or '').strip() for v in next(iterator)]
    except StopIteration: return []
    rows=[]
    for values in iterator:
        if all(v is None or str(v).strip()=='' for v in values): continue
        rows.append({header[i]:values[i] if i < len(values) else None for i in range(len(header))})
    wb.close(); return rows


def read_tabular(path: str | Path) -> list[dict[str,Any]]:
    p=Path(path); ext=p.suffix.lower()
    if ext=='.csv': return _read_csv(p)
    if ext in {'.xlsx','.xlsm'}: return _read_xlsx(p)
    raise ValueError('Поддерживаются только CSV и XLSX.')


def import_costs(repo: Repository, shop_id: int, path: str | Path, *,
                 default_effective_date: str | None = None) -> ImportSummary:
    p=Path(path); rows=read_tabular(p); batch=repo.create_import_batch(shop_id,'costs',p.name)
    errors: list[dict[str,Any]]=[]; applied=0
    default_day=default_effective_date or date.today().isoformat()
    for num,raw in enumerate(rows,start=2):
        row=_headers(raw)
        try:
            market=str(row.get('marketplace') or row.get('market') or '').strip().lower()
            market={'wb':'wildberries','wildberries':'wildberries','ozon':'ozon'}.get(market,'')
            sku=str(row.get('sku') or row.get('marketplace_sku') or row.get('nmid') or '').strip()
            if not market or not sku: raise ValueError('нужны marketplace и sku')
            value=row.get('cost_price') if row.get('cost_price') is not None else row.get('cost')
            cost=float(str(value).replace(' ','').replace(',','.'))
            if cost < 0: raise ValueError('cost_price не может быть отрицательной')
            day=str(row.get('effective_date') or default_day)[:10]
            listing=repo.listing_for_shop(shop_id,market,sku)
            if not listing: raise ValueError('SKU не найден; сначала выполните /backfill')
            internal=str(row.get('internal_sku') or '').strip()
            name=str(row.get('name') or '').strip() or None
            if internal:
                product=repo.link_marketplace_listings(shop_id,internal,[(market,sku)],name=name)
                product_id=product.id
            else:
                product_id=int(listing['product_id'])
            if not repo.set_product_cost(product_id,cost,effective_date=day,
                                         source=f'import:{p.name}',import_batch_id=batch):
                raise ValueError('товар не найден')
            applied += 1
        except Exception as exc:
            errors.append({'row':num,'error':str(exc)[:300]})
    status='success' if not errors else ('partial' if applied else 'failed')
    repo.finish_import_batch(batch,status=status,total_rows=len(rows),applied_rows=applied,errors=errors)
    return ImportSummary(batch,p.name,len(rows),applied,tuple(errors))
