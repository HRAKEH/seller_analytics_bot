from datetime import date
from pathlib import Path
import zipfile

from openpyxl import load_workbook

from app.config import Settings
from app.services.backups import BackupService, inspect_database
from app.services.exporting import export_xlsx, export_csv_zip
from app.storage import Database, Repository, MetricPoint, LATEST_SCHEMA_VERSION
from app.storage.models import ProductMetricPoint


def make_repo(tmp_path: Path):
    db=Database(tmp_path/'phase10.sqlite3'); assert db.initialize()==LATEST_SCHEMA_VERSION
    repo=Repository(db); seller=repo.ensure_seller(1,'Seller'); shop=repo.ensure_shop(seller.id,'Shop A')
    return db,repo,seller,shop


def test_per_user_shop_selection_and_profiles(tmp_path):
    _,repo,seller,a=make_repo(tmp_path)
    b=repo.ensure_shop(seller.id,'Shop B',credential_profile='SHOP2')
    assert b.credential_profile=='SHOP2'
    repo.select_shop_for_user(100,a.id); repo.select_shop_for_user(200,b.id)
    assert repo.selected_shop_for_user(100)==a.id
    assert repo.selected_shop_for_user(200)==b.id
    assert {x.id for x in repo.list_shops(seller.id)}=={a.id,b.id}
    repo.set_shop_credential_profile(b.id,'SECOND')
    assert repo.get_shop(b.id).credential_profile=='SECOND'


def test_profile_credentials_are_resolved_from_env(monkeypatch):
    monkeypatch.setenv('SELLERBOT_STORE2_WB_API_TOKEN','wb-secret')
    monkeypatch.setenv('SELLERBOT_STORE2_OZON_CLIENT_ID','cid')
    monkeypatch.setenv('SELLERBOT_STORE2_OZON_API_KEY','key')
    settings=Settings.from_env(); creds=settings.credentials_for_profile('store2')
    assert creds.has_wb and creds.has_ozon
    assert creds.wb_api_token=='wb-secret' and creds.ozon_client_id=='cid'
    assert 'STORE2' in settings.available_credential_profiles()


def test_backup_restore_rolls_database_back(tmp_path):
    db,repo,_,shop=make_repo(tmp_path)
    service=BackupService(db,repo,tmp_path/'backups')
    backup=service.create(kind='manual')
    version,ok=inspect_database(backup.path)
    assert version==LATEST_SCHEMA_VERSION and ok and backup.size_bytes>0
    repo.rename_shop(shop.id,'Changed')
    assert repo.get_shop(shop.id).name=='Changed'
    service.restore(backup.path)
    assert repo.get_shop(shop.id).name=='Shop A'
    assert db.integrity_check()


def test_export_xlsx_and_csv_zip(tmp_path):
    db,repo,_,shop=make_repo(tmp_path)
    conn=repo.ensure_connection(shop.id,'ozon','Ozon')
    product=repo.ensure_product(shop.id,'P1','Product',cost_price=100)
    listing=repo.ensure_listing(product.id,conn.id,'1001','O-1001')
    run=repo.record_success(conn.id,'analytics/orders','2026-09-28',{'x':1},[
        MetricPoint(conn.id,'2026-09-28','ordered_units',2,'units',True),
        MetricPoint(conn.id,'2026-09-28','ordered_revenue',500,'RUB',True),
    ])
    repo.save_product_metrics(run,[
        ProductMetricPoint(listing.id,'2026-09-28','ordered_units',2,'units'),
        ProductMetricPoint(listing.id,'2026-09-28','ordered_revenue',500,'RUB'),
    ])
    xlsx=export_xlsx(repo,shop.id,date(2026,9,28),1,tmp_path/'export.xlsx')
    wb=load_workbook(xlsx.path,read_only=True)
    assert {'О файле','Показатели по дням','Товары','Остатки','Финансовая сводка','Сверка данных','Себестоимость','Описание полей'} <= set(wb.sheetnames)
    assert wb['Товары'].max_row>=2
    csv_result=export_csv_zip(repo,shop.id,date(2026,9,28),1,tmp_path/'export_csv.zip')
    with zipfile.ZipFile(csv_result.path) as z:
        names=set(z.namelist())
        assert 'summary.csv' in names and 'products.csv' in names and 'daily.csv' in names
        assert b'P1' in z.read('products.csv')


def test_job_state_is_persistent(tmp_path):
    _,repo,_,shop=make_repo(tmp_path)
    assert repo.get_job_state(shop.id,'daily_report') is None
    repo.set_job_state(shop.id,'daily_report','2026-09-29')
    assert repo.get_job_state(shop.id,'daily_report')=='2026-09-29'


def test_schema_v5_migrates_to_latest(tmp_path):
    import sqlite3
    from app.storage.database import (_migration_1,_migration_2,_migration_3,_migration_4,_migration_5)
    path=tmp_path/'v5.sqlite3'
    conn=sqlite3.connect(path)
    conn.execute('CREATE TABLE schema_migrations (version INTEGER PRIMARY KEY, applied_at TEXT NOT NULL)')
    for version,fn in ((1,_migration_1),(2,_migration_2),(3,_migration_3),(4,_migration_4),(5,_migration_5)):
        fn(conn); conn.execute("INSERT INTO schema_migrations VALUES(?,datetime('now'))",(version,))
    conn.commit(); conn.close()
    db=Database(path)
    assert db.initialize()==LATEST_SCHEMA_VERSION
    with db.connect() as c:
        columns={r[1] for r in c.execute('PRAGMA table_info(shops)')}
        tables={r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    assert 'credential_profile' in columns
    assert {'user_shop_selection','shop_job_state','backup_history'} <= tables
    assert db.integrity_check()
