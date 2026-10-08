"""Real Telegram handlers, persisted sources and files; no external API calls."""
import csv
import io
from datetime import date, datetime
from pathlib import Path
from unittest.mock import AsyncMock
import zipfile

from aiogram.methods import SendDocument, SendMessage
from openpyxl import load_workbook, Workbook
import pytest

from test_navigation import ui, press, callback, active, inline_data
from app.bot.keyboards import COMMAND_BUTTONS
from app.reports.accruals import build_accrual_ledger, export_accrual_ledger
from app.reports.wb_accruals import build_wb_accrual_ledger
from app.reports.accrual_cards import build_accrual_card, accrual_totals
from app.services.finance import normalize_ozon_accruals
from app.services.imports import import_costs
from app.services.spreadsheet_format import write_sheet
from app.storage import Database, Repository, LATEST_SCHEMA_VERSION, MetricPoint
from app.storage.database import MIGRATIONS

DAY = '2026-09-28'


def seed_ozon(ui, *, shop=None, amount='80.15'):
    shop = shop or ui.shop
    conn = ui.repo.ensure_connection(shop.id, 'ozon', 'Ozon')
    payload = {'accruals': [
        {'id': '001234567890123456789', 'unit_number': '0007',
         'date': DAY+'T09:12:00Z', 'total_amount': amount,
         'posting': {'posting_number': '00001-0002', 'products': [
             {'sku': '000123', 'commission': {'seller_price': '100.25'}}]},
         'non_item_fee': {'type_id': 54, 'accrued': '-20.10'}},
        {'id': '=BAD', 'total_amount': '2.05', 'non_item_fee': {'type_id': 54, 'accrued': '2.05'}},
    ]}
    run = ui.repo.record_success(conn.id, 'finance/accrual/by-day', DAY, payload,
                                normalize_ozon_accruals(payload, conn.id, DAY))
    return conn, run


def seed_wb(ui, *, count=1):
    conn = ui.repo.ensure_connection(ui.shop.id, 'wildberries', 'WB')
    reports = [{'reportId': 1000+i, 'dateFrom': '2026-09-21', 'dateTo': '2026-09-27',
                'retailAmountSum': 100.25, 'forPaySum': 80.15, 'bankPaymentSum': 70.05}
               for i in range(count)]
    run = ui.repo.record_success(conn.id, 'finance/sales-reports/list', DAY, {'reports': reports}, [])
    return conn, run


def last_card(ui, *, user=101):
    return max(mid for (chat, mid), msg in ui.telegram.messages.items()
               if chat == user and msg.from_user.is_bot and msg.text and '🧾' in msg.text)


def capture_downloads(ui):
    files = []
    original = ui.telegram.request

    async def request(bot, method, **kwargs):
        if isinstance(method, SendDocument):
            path = Path(method.document.path)
            files.append((path.name, path.read_bytes(), method.caption))
        return await original(bot, method, **kwargs)

    ui.bot.session.make_request = AsyncMock(side_effect=request)
    return files


@pytest.mark.asyncio
@pytest.mark.parametrize('market', ['ozon', 'wb'])
async def test_accruals_open_as_one_card_expand_collapse_and_download_both_formats(ui, market):
    seed_ozon(ui) if market == 'ozon' else seed_wb(ui)
    files = capture_downloads(ui)
    command = '/accruals' if market == 'ozon' else '/wb_accruals'
    await press(ui, command+' 1 '+DAY)
    card = last_card(ui)
    brief = ui.telegram.messages[(101, card)].text
    assert not files
    assert '28.09.2026' in brief and 'Shop' in brief
    if market == 'wb':
        assert '21.09.2026' in brief and '27.09.2026' in brief
        assert 'Пропуски не считаются нулём' in brief
    await callback(ui, card, inline_data(ui, card, 'Подробнее'))
    assert ui.telegram.messages[(101, card)].text != brief
    await callback(ui, card, inline_data(ui, card, 'Свернуть'))
    assert ui.telegram.messages[(101, card)].text == brief
    await callback(ui, card, inline_data(ui, card, '📄 Скачать CSV'))
    await callback(ui, card, inline_data(ui, card, '📊 Excel (.xlsx)'))
    assert len(files) == 2 and files[0][0].endswith('.csv') and files[1][0].endswith('.xlsx')
    csv_rows = list(csv.DictReader(io.StringIO(files[0][1].decode('utf-8-sig')), delimiter=';'))
    wb = load_workbook(io.BytesIO(files[1][1]))
    ws = wb['Операции Ozon' if market == 'ozon' else 'Отчёты WB']
    assert ws.freeze_panes == 'A2' and ws.auto_filter.ref
    values = list(ws.values)
    assert len(values)-1 == len(csv_rows)
    key = 'marketplace_net' if market == 'ozon' else 'bank_payment'
    column = values[0].index('Начисления после удержаний, ₽' if market == 'ozon' else 'К оплате по отчёту, ₽')
    assert sum(float(row[key]) for row in csv_rows) == pytest.approx(sum(row[column] for row in values[1:]))
    if market == 'ozon':
        assert values[1][values[0].index('Номер операции')] == '001234567890123456789'
        assert values[1][values[0].index('Артикулы')] == '000123'
        assert values[2][values[0].index('Номер операции')] == "'=BAD"
        assert values[2][values[0].index('Услуги и прочие удержания, ₽')] < 0
    else:
        assert '21-09-2026_27-09-2026' in files[1][0]
        assert values[1][values[0].index('Хранение, ₽')] is None
        summary = dict(wb['Итог'].values)
        assert summary['Даты сохранения с'] == datetime(2026, 9, 28)
        assert summary['Период отчёта WB с'] == datetime(2026, 9, 21)


@pytest.mark.asyncio
async def test_card_download_keeps_original_source_after_correction_and_restart(ui):
    _, old_run = seed_ozon(ui)
    files = capture_downloads(ui)
    await press(ui, '/accruals 1 '+DAY)
    old_card = last_card(ui)
    assert ui.repo.data_view(ui.bot.id, 101, old_card)['payload']['source_run_ids'] == [old_run]
    seed_ozon(ui, amount='800.15')
    ui.dp = ui.new_dispatcher()
    await callback(ui, old_card, inline_data(ui, old_card, '📄 Скачать CSV'))
    rows = list(csv.DictReader(io.StringIO(files[-1][1].decode('utf-8-sig')), delimiter=';'))
    assert sum(float(row['marketplace_net']) for row in rows) == pytest.approx(82.2)
    await press(ui, '/accruals 1 '+DAY)
    assert '802,20' in ui.telegram.messages[(101, last_card(ui))].text


@pytest.mark.asyncio
async def test_wb_long_details_page_in_place_survives_restart_and_keeps_warning(ui):
    seed_wb(ui, count=40)
    await press(ui, '/wb_accruals 1 '+DAY)
    card = last_card(ui)
    await callback(ui, card, inline_data(ui, card, 'Подробнее'))
    before = ui.telegram.messages[(101, card)].text
    assert 'Страница 1/' in before and len(before) < 4096
    ui.dp = ui.new_dispatcher()
    await callback(ui, card, inline_data(ui, card, '▶️'))
    after = ui.telegram.messages[(101, card)].text
    assert 'Страница 2/' in after and before != after
    assert 'Пропуски не считаются нулём' in after
    view = ui.repo.data_view(ui.bot.id, 101, card)
    assert '<code>1000</code>' in ''.join(view['payload']['pages'])
    assert len([m for m in ui.telegram.methods if isinstance(m, SendMessage)]) == 1


@pytest.mark.asyncio
@pytest.mark.parametrize('restriction', ['other_user', 'manager', 'revoked', 'archived'])
async def test_finance_card_download_checks_current_access_and_original_user(ui, restriction):
    seed_ozon(ui)
    files = capture_downloads(ui)
    await press(ui, '/accruals 1 '+DAY)
    card = last_card(ui)
    user = 101
    if restriction == 'other_user':
        user = 102
        ui.telegram.messages[(102, card)] = ui.telegram.messages[(101, card)]
    elif restriction == 'manager':
        ui.repo.grant_shop_access(101, ui.shop.id, 'manager')
    elif restriction == 'revoked':
        ui.repo.revoke_shop_access(101, ui.shop.id)
    else:
        ui.repo.ensure_shop(ui.shop.seller_id, 'Other')
        ui.repo.archive_shop(101, ui.shop.id)
    await callback(ui, card, 'data_view:file:xlsx', user=user)
    assert not files
    assert 'доступ' in ui.telegram.methods[-1].text.lower()


@pytest.mark.asyncio
async def test_shop_switch_does_not_change_original_card_or_export(ui):
    seed_ozon(ui)
    other = ui.repo.ensure_shop(ui.shop.seller_id, 'Другой магазин')
    ui.repo.grant_shop_access(101, other.id, 'owner')
    seed_ozon(ui, shop=other, amount='9999')
    files = capture_downloads(ui)
    await press(ui, '/accruals 1 '+DAY)
    card = last_card(ui)
    await press(ui, '/export')
    picker = active(ui)
    ui.ctx.shop_id = other.id
    ui.dp = ui.new_dispatcher()
    await callback(ui, card, inline_data(ui, card, '📄 Скачать CSV'))
    await callback(ui, picker, inline_data(ui, picker, '7 дней'))
    await callback(ui, picker, inline_data(ui, picker, '📊 Скачать Excel (.xlsx)'))
    assert all('Shop' in caption and 'Другой магазин' not in caption for _, _, caption in files)
    assert all('Shop' in name for name, _, _ in files)
    assert build_accrual_ledger(ui.repo, ui.shop.id, DAY, DAY).rows[0]['marketplace_net'] == 80.15


@pytest.mark.asyncio
async def test_export_menu_uses_buttons_pages_and_current_role_with_safe_metadata(ui):
    ui.repo.grant_shop_access(102, ui.shop.id, 'accountant')
    conn, _ = seed_ozon(ui)
    ui.repo.record_success(conn.id, 'analytics/orders', DAY, {}, [
        MetricPoint(conn.id, DAY, 'ordered_units', 3, 'units'),
        MetricPoint(conn.id, DAY, 'commission', 987654, 'RUB')])
    files = capture_downloads(ui)
    await press(ui, COMMAND_BUTTONS['export'], user=102)
    picker = active(ui, 102)
    assert 'Выберите период' in ui.telegram.messages[(102, picker)].text
    await callback(ui, picker, inline_data(ui, picker, '7 дней', user=102), user=102)
    await callback(ui, picker, inline_data(ui, picker, 'Что будет в файле', user=102), user=102)
    await callback(ui, picker, inline_data(ui, picker, '▶️', user=102), user=102)
    assert 'Страница 2/' in ui.telegram.messages[(102, picker)].text
    ui.repo.grant_shop_access(102, ui.shop.id, 'manager')
    ui.dp = ui.new_dispatcher()
    await callback(ui, picker, 'data_view:file:csv', user=102)
    with zipfile.ZipFile(io.BytesIO(files[-1][1])) as archive:
        names = archive.namelist()
        assert {'Описание_файлов.txt', 'Описание_полей.csv'} <= set(names)
        assert 'finance.csv' not in names and 'costs.csv' not in names
        content = '\n'.join(archive.read(name).decode('utf-8-sig') for name in names)
        assert 'cost_price' not in content and 'credential_profile' not in content
    assert active(ui, 102) == picker


@pytest.mark.asyncio
async def test_export_rechecks_finance_rights_after_file_generation(ui, monkeypatch):
    ui.repo.grant_shop_access(102, ui.shop.id, 'accountant')
    files = capture_downloads(ui)
    await press(ui, '/export', user=102)
    picker = active(ui, 102)
    await callback(ui, picker, 'data_view:days:7', user=102)

    def exporter(repo, shop_id, end, days, path, **permissions):
        assert permissions['include_finance']
        path.write_bytes(b'PRIVATE FINANCIAL DATA')
        repo.grant_shop_access(102, shop_id, 'manager')

    monkeypatch.setattr('app.bot.data_views.export_xlsx', exporter)
    await callback(ui, picker, 'data_view:file:xlsx', user=102)
    assert not files and 'Доступ изменился' in ui.telegram.methods[-1].text


@pytest.mark.asyncio
async def test_legacy_export_callback_does_not_download_from_current_shop(ui):
    files = capture_downloads(ui)
    await press(ui, '/export')
    picker = active(ui)
    await callback(ui, picker, 'export:run:30:xlsx')
    assert not files and 'заново' in ui.telegram.methods[-1].text


def test_ledgers_pin_only_sources_from_the_selected_shop_and_missing_source_fails(ui):
    _, run = seed_ozon(ui)
    other = ui.repo.ensure_shop(ui.shop.seller_id, 'Other')
    _, other_run = seed_ozon(ui, shop=other)
    with pytest.raises(ValueError):
        build_accrual_ledger(ui.repo, ui.shop.id, DAY, DAY, source_run_ids=[other_run])
    with ui.repo.db.connect() as c:
        c.execute('DELETE FROM raw_payloads WHERE source_run_id=?', (run,))
    with pytest.raises(ValueError):
        build_accrual_ledger(ui.repo, ui.shop.id, DAY, DAY, source_run_ids=[run])


def test_known_empty_day_is_zero_missing_day_is_unknown_and_has_no_download(ui):
    conn = ui.repo.ensure_connection(ui.shop.id, 'ozon', 'Ozon')
    ui.repo.record_success(conn.id, 'finance/accrual/by-day', DAY, {'accruals': []}, [])
    empty = build_accrual_ledger(ui.repo, ui.shop.id, DAY, DAY)
    missing = build_accrual_ledger(ui.repo, ui.shop.id, '2026-09-29', '2026-09-29')
    assert accrual_totals(empty, 'ozon')['marketplace_net'] == 0
    assert accrual_totals(missing, 'ozon')['marketplace_net'] is None
    assert build_accrual_card(empty, 'ozon', 'Shop')['available']
    assert not build_accrual_card(missing, 'ozon', 'Shop')['available']


def test_excel_dates_timezone_currency_and_cost_import_preserve_values(ui, tmp_path):
    conn, _ = seed_ozon(ui)
    product = ui.repo.ensure_product(ui.shop.id, 'P1', '=BAD', cost_price=12.34)
    ui.repo.ensure_listing(product.id, conn.id, '000123')
    wb = Workbook(); wb.remove(wb.active)
    ws = write_sheet(wb, 'Summary', [
        {'key': 'period_start', 'value': DAY},
        {'key': 'generated_at_utc', 'value': DAY+'T23:09:00Z'}], timezone='Europe/Moscow')
    assert ws['B2'].value == date(2026, 9, 28) and ws['B2'].number_format == 'dd.mm.yyyy'
    assert ws['B3'].value == datetime(2026, 9, 29, 2, 9) and ws['B3'].number_format == 'dd.mm.yyyy hh:mm'
    daily = write_sheet(wb, 'Daily', [{'metric_key': 'ordered_revenue', 'value': 12.34, 'unit': 'RUB'}])
    assert '₽' in daily['B2'].number_format
    costs = Workbook(); costs.remove(costs.active)
    write_sheet(costs, 'Costs', [{'marketplace': 'ozon', 'marketplace_sku': '000123',
        'cost_price': 23.45, 'effective_date': DAY}])
    path = tmp_path/'costs.xlsx'; costs.save(path)
    result = import_costs(ui.repo, ui.shop.id, path)
    assert result.applied_rows == 1 and not result.errors
    assert ui.repo.listing_for_shop(ui.shop.id, 'ozon', '000123')['cost_price'] == 23.45


def test_schema20_upgrade_preserves_records_and_binds_card_identity(tmp_path):
    db = Database(tmp_path/'old.sqlite3')
    with db.connect() as c:
        c.execute('CREATE TABLE schema_migrations (version INTEGER PRIMARY KEY, applied_at TEXT NOT NULL)')
        for version, migration in MIGRATIONS.items():
            if version > 20: break
            migration(c)
            c.execute("INSERT INTO schema_migrations VALUES(?,datetime('now'))", (version,))
    repo = Repository(db)
    shop = repo.ensure_shop(repo.ensure_seller(101).id, 'Preserved')
    repo.grant_shop_access(101, shop.id, 'owner')
    conn = repo.ensure_connection(shop.id, 'ozon', 'Ozon')
    run = repo.record_success(conn.id, 'finance/accrual/by-day', DAY, {'accruals': []}, [])
    assert db.initialize_safely(tmp_path/'backups') == LATEST_SCHEMA_VERSION
    assert db.integrity_check() and repo.get_shop(shop.id).name == 'Preserved'
    assert build_accrual_ledger(repo, shop.id, DAY, DAY).source_run_ids == (run,)
    assert list((tmp_path/'backups').glob('*.sqlite3'))
    repo.save_data_view(123, 101, 9, shop_id=shop.id, user_id=101, kind='ozon', payload={'page': 0})
    repo.save_data_view(123, 101, 9, shop_id=shop.id, user_id=102, kind='export', payload={'page': 9})
    saved = repo.data_view(123, 101, 9)
    assert saved['user_id'] == 101 and saved['kind'] == 'ozon' and saved['payload'] == {'page': 0}
