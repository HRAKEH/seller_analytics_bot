"""Approved shop roles exercised through real handlers and durable reports."""
from __future__ import annotations

import asyncio
from concurrent.futures import ThreadPoolExecutor
from contextlib import asynccontextmanager
from datetime import date
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock
import zipfile

from aiogram import Dispatcher, types
from aiogram.methods import AnswerCallbackQuery, SendDocument, SendMessage
from openpyxl import load_workbook
import pytest

from app.access import PERMISSIONS, can_role
from app.bot.handlers import register_handlers
from app.bot.handlers import CostImportStates
from app.bot.keyboards import COMMAND_BUTTONS, HOME, CANCEL, MENU_TECH, action_ref
from app.bot.paged_reports import PagedReportController
from app.bot.report_cards import send_daily_card
from app.bot.runtime import ShopContextMiddleware
from app.reports.cards import DailyCardText
from app.services.actions import ActionCenter, ActionItem
from app.services.exporting import collect_export_tables, export_xlsx, export_csv_zip
from app.services.report_refresh import refresh_reports
from app.storage import Database, Repository, MetricPoint, ProductMetricPoint

from test_navigation import ui, press, callback, active, inline_data
from test_report_cards import tap, stored


def last_text(ui,user=101):
    return next(m.text for m in reversed(ui.telegram.methods) if isinstance(m,SendMessage) and m.chat_id==user)


def buttons(message):
    return [b.text for row in message.reply_markup.inline_keyboard for b in row] if message.reply_markup else []


@pytest.mark.parametrize('role,expected',[
    ('owner',PERMISSIONS),('accountant',{'view','operate','finance','costs','settings'}),
    ('manager',{'view','operate'}),('viewer',{'view','finance'}),('analyst',{'view','operate','finance','costs','settings'}),
    (None,set())])
def test_approved_permissions(role,expected):
    assert {permission for permission in PERMISSIONS if can_role(role,permission)}==set(expected)


def test_atomic_editor_checks_actor_shop_and_keeps_other_access(ui):
    other=ui.repo.ensure_shop(ui.shop.seller_id,'Other')
    ui.repo.grant_shop_access(101,other.id,'owner')
    ui.repo.manage_employee_access(101,ui.shop.id,103,'manager',display_name='Employee')
    ui.repo.manage_employee_access(101,other.id,103,'accountant')
    assert ui.repo.role_for_user(103,ui.shop.id)=='manager'
    assert ui.repo.role_for_user(103,other.id)=='accountant'
    for actor in (103,999):
        with pytest.raises(PermissionError):ui.repo.manage_employee_access(actor,ui.shop.id,104,'owner')
    ui.repo.manage_employee_access(101,ui.shop.id,103,None)
    assert ui.repo.role_for_user(103,ui.shop.id) is None
    assert ui.repo.role_for_user(103,other.id)=='accountant'


def test_editor_protects_last_owner_protected_owner_and_stale_role(ui):
    for role in (None,'manager','accountant'):
        with pytest.raises(ValueError,match='последнего владельца'):
            ui.repo.manage_employee_access(101,ui.shop.id,101,role)
    ui.repo.grant_shop_access(103,ui.shop.id,'owner')
    with pytest.raises(ValueError,match='владелец всего бота'):
        ui.repo.manage_employee_access(103,ui.shop.id,101,'manager',protected_user_ids=(101,))
    ui.repo.manage_employee_access(101,ui.shop.id,104,'accountant')
    with pytest.raises(ValueError,match='уже изменился'):
        ui.repo.manage_employee_access(101,ui.shop.id,104,'manager',expected_role=None)
    assert ui.repo.role_for_user(104,ui.shop.id)=='accountant'


def test_concurrent_owner_demotions_never_remove_last_owner(ui):
    ui.repo.grant_shop_access(103,ui.shop.id,'owner')
    def demote(uid):
        try:ui.repo.manage_employee_access(uid,ui.shop.id,uid,'manager');return True
        except (PermissionError,ValueError):return False
    with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(demote,(101,103)))
    assert results.count(True)==1
    assert len(ui.repo.users_for_shop(ui.shop.id,roles=['owner']))==1


@pytest.mark.parametrize('uid',[0,-1,2**63])
def test_editor_rejects_invalid_ids(ui,uid):
    with pytest.raises(ValueError):ui.repo.manage_employee_access(101,ui.shop.id,uid,'manager')


@pytest.mark.asyncio
async def test_button_employee_add_edit_revoke_in_same_message(ui):
    await press(ui,COMMAND_BUTTONS['users']);mid=active(ui)
    await callback(ui,mid,inline_data(ui,mid,'➕ Добавить сотрудника'))
    await press(ui,'103')
    assert active(ui)==mid
    assert {'Владелец','Бухгалтер','Менеджер'}<=set(buttons(ui.telegram.messages[(101,mid)]))
    await callback(ui,mid,inline_data(ui,mid,'Менеджер'))
    assert ui.repo.role_for_user(103,ui.shop.id)=='manager'
    assert active(ui)==mid
    await callback(ui,mid,inline_data(ui,mid,'103 · Менеджер'))
    await callback(ui,mid,inline_data(ui,mid,'✏️ Изменить роль'))
    await callback(ui,mid,inline_data(ui,mid,'Бухгалтер'))
    assert ui.repo.role_for_user(103,ui.shop.id)=='accountant'
    await callback(ui,mid,inline_data(ui,mid,'103 · Бухгалтер'))
    await callback(ui,mid,inline_data(ui,mid,'➖ Отозвать доступ'))
    assert ui.repo.role_for_user(103,ui.shop.id)=='accountant'
    await callback(ui,mid,inline_data(ui,mid,'➖ Да, отозвать'))
    assert ui.repo.role_for_user(103,ui.shop.id) is None
    assert active(ui)==mid


@pytest.mark.asyncio
@pytest.mark.parametrize('navigation',[HOME,CANCEL,'/start','/cancel'])
async def test_employee_id_prompt_does_not_consume_navigation(ui,navigation):
    await press(ui,COMMAND_BUTTONS['user_add'])
    await press(ui,navigation)
    assert 'числовой Telegram ID, без' not in last_text(ui)
    assert not any(row['telegram_user_id']==103 for row in ui.repo.users_for_shop(ui.shop.id))


@pytest.mark.asyncio
async def test_employee_form_rejects_stale_other_actor_and_revoked_owner(ui):
    await press(ui,COMMAND_BUTTONS['user_add']);mid=active(ui)
    await press(ui,'103');choose=inline_data(ui,mid,'Менеджер')
    # A callback carrying another person's old message must not grant access.
    ui.telegram.messages[(102,mid)]=ui.telegram.messages[(101,mid)]
    await callback(ui,mid,choose,user=102)
    assert ui.repo.role_for_user(103,ui.shop.id) is None
    ui.repo.grant_shop_access(103,ui.shop.id,'accountant')
    await callback(ui,mid,choose)
    assert ui.repo.role_for_user(103,ui.shop.id)=='accountant'
    ui.repo.grant_shop_access(101,ui.shop.id,'accountant')
    await callback(ui,mid,choose)
    assert ui.repo.role_for_user(103,ui.shop.id)=='accountant'


@pytest.mark.asyncio
async def test_employee_list_has_pages_and_copyable_ids(ui):
    for uid in range(103,113):ui.repo.grant_shop_access(uid,ui.shop.id,'manager',f'Employee {uid}')
    await press(ui,'/users');mid=active(ui)
    assert len([b for b in buttons(ui.telegram.messages[(101,mid)]) if 'Менеджер' in b or 'Владелец' in b])<=5
    assert '<code>101</code>' in ui.telegram.messages[(101,mid)].text
    await callback(ui,mid,inline_data(ui,mid,'▶️'))
    assert '2/3' in buttons(ui.telegram.messages[(101,mid)])


CLOSED_FINANCE=['/finance','/management','/sku_finance','/reconcile','/accruals','/wb_accruals','/sources','/readiness',
    '/cost wb 123 99','/import_costs','/link P wb 123','🔎 Сверка','📈 Результат',COMMAND_BUTTONS['finance'],
    '🔄 Обновить финансы',COMMAND_BUTTONS['cost'],COMMAND_BUTTONS['import_costs']]


@pytest.mark.asyncio
@pytest.mark.parametrize('command',CLOSED_FINANCE+['/setup','/supply_defaults 1 2 3','/supply_set P 1 2 3',
    '/users','/user_add 104 owner','/user_remove 101'])
async def test_manager_closed_commands_and_old_buttons_do_not_execute(ui,command):
    ui.repo.grant_shop_access(103,ui.shop.id,'manager')
    ui.ctx.collect_finance=AsyncMock(side_effect=AssertionError('closed finance'))
    ui.ctx.collect_reconciliation=AsyncMock(side_effect=AssertionError('closed reconciliation'))
    before=ui.repo.users_for_shop(ui.shop.id)
    await press(ui,command,user=103)
    assert any(word in last_text(ui,103) for word in ('Недостаточно прав','только владелец'))
    assert ui.repo.users_for_shop(ui.shop.id)==before
    ui.ctx.collect_finance.assert_not_called();ui.ctx.collect_reconciliation.assert_not_called()


@pytest.mark.asyncio
@pytest.mark.parametrize('role',['accountant','manager','owner'])
@pytest.mark.parametrize('command',['/backup','/backups','/restore','/health','/jobs','/diagnostics','/profiles','/help',
    '/connect_check','/demo_on','/demo_off','/shop_add','/shop_profile','/shop_archive','/shop_archived','/shop_restore','/shop_delete',MENU_TECH])
async def test_technical_commands_require_installation_owner(ui,role,command):
    ui.repo.grant_shop_access(103,ui.shop.id,role)
    await press(ui,command,user=103)
    assert 'Недостаточно прав' in last_text(ui,103)


@pytest.mark.asyncio
async def test_accountant_can_change_working_settings_and_costs_but_not_employees(ui):
    ui.repo.grant_shop_access(103,ui.shop.id,'accountant')
    conn=ui.repo.ensure_connection(ui.shop.id,'wildberries','WB')
    product=ui.repo.ensure_product(ui.shop.id,'P','Product');ui.repo.ensure_listing(product.id,conn.id,'123')
    await press(ui,'/cost wb 123 99 2026-09-28',user=103)
    assert '✅' in last_text(ui,103)
    await press(ui,'/supply_defaults 3 4 5',user=103)
    assert ui.repo.ensure_shop_supply_preferences(ui.shop.id)['default_lead_time_days']==3
    await press(ui,'/setup',user=103)
    await press(ui,'Accounting Shop',user=103)
    assert ui.repo.get_shop(ui.shop.id).name=='Accounting Shop'
    await press(ui,CANCEL,user=103)
    await press(ui,'/users',user=103)
    assert 'Недостаточно прав' in last_text(ui,103)


@pytest.mark.asyncio
async def test_role_change_during_settings_input_prevents_mutation(ui):
    ui.repo.grant_shop_access(103,ui.shop.id,'accountant')
    await press(ui,'/setup',user=103)
    ui.repo.grant_shop_access(103,ui.shop.id,'manager')
    await press(ui,'Must not save',user=103)
    assert ui.repo.get_shop(ui.shop.id).name=='Shop'
    assert 'Недостаточно прав' in last_text(ui,103)


@pytest.mark.asyncio
async def test_manager_daily_card_hides_accruals_and_old_callbacks_after_demotion(ui):
    ui.repo.grant_shop_access(103,ui.shop.id,'accountant')
    text=DailyCardText('Public summary','Public detail','Private finance 987654,32 ₽')
    card=await send_daily_card(ui.bot,ui.ctx,103,date(2026,10,1),user_id=103,text=text)
    await tap(ui,card,'accruals',user=103)
    ui.repo.grant_shop_access(103,ui.shop.id,'manager')
    before=len(ui.telegram.methods)
    await tap(ui,card,'page:0',user=103)
    assert [type(m).__name__ for m in ui.telegram.methods[before:]]==['AnswerCallbackQuery']
    await tap(ui,card,'collapse',user=103)
    current=ui.telegram.messages[(103,card.message_id)]
    assert 'Private finance' not in current.text and 'Начисления' not in buttons(current)
    await tap(ui,card,'accruals',user=103)
    assert stored(ui,card)['section']=='summary'
    ui.ctx.refresh_reports=AsyncMock(return_value=[])
    await tap(ui,card,'refresh',user=103)
    assert ui.ctx.refresh_reports.await_args.kwargs=={'include_finance':False}


@pytest.mark.asyncio
async def test_manager_refresh_excludes_finance_and_does_not_print_financial_result(ui):
    ui.repo.grant_shop_access(103,ui.shop.id,'manager')
    ui.ctx.refresh_reports=AsyncMock(return_value=[])
    await press(ui,'/refresh 1 2026-09-28',user=103)
    assert ui.ctx.refresh_reports.await_args.kwargs['include_finance'] is False
    assert not any('Финансовые итоги' in str(getattr(m,'text','')) for m in ui.telegram.methods)


@pytest.mark.asyncio
async def test_refresh_service_skips_finance_sources():
    @asynccontextmanager
    async def lock(name):yield
    result=[SimpleNamespace(ok=True)]
    collector=SimpleNamespace(collect_wb_orders_day=AsyncMock(return_value=result[0]),
        collect_ozon_orders_day=AsyncMock(return_value=result[0]),collect_finance=AsyncMock(return_value=result),
        collect_advertising=AsyncMock(return_value=result),collect_ozon_fulfillment_range=AsyncMock(return_value=result),
        collect_ozon_buyout_prices_range=AsyncMock(return_value=result),collect_wb_sales_range=AsyncMock(return_value=result))
    ctx=SimpleNamespace(demo_mode=lambda:False,operation_lock=lock,shop_id=1,collector=collector,wb_connection_id=2,ozon_connection_id=3)
    stages=await refresh_reports(ctx,date(2026,9,28),date(2026,9,28),include_finance=False)
    collector.collect_finance.assert_not_called()
    assert not any(s.key.startswith('finance') for s in stages)
    assert collector.collect_advertising.await_count==2


@pytest.mark.asyncio
async def test_finance_paged_snapshot_rechecks_new_capability_after_restart(ui):
    ui.repo.grant_shop_access(103,ui.shop.id,'accountant')
    await press(ui,'/my_access',user=103)
    message=ui.telegram.messages[(103,ui.update_id)].as_(ui.bot)
    sent=await PagedReportController(ui.ctx).show(message,'Private finance\n'+'a'*5000,permission='finance')
    assert ui.repo.paged_report(ui.bot.id,103,sent.message_id)['capability']=='finance'
    ui.repo.grant_shop_access(103,ui.shop.id,'manager')
    ui.dp=ui.new_dispatcher();before=len(ui.telegram.methods)
    await callback(ui,sent.message_id,'report_page:1',user=103)
    assert [type(m).__name__ for m in ui.telegram.methods[before:]]==['AnswerCallbackQuery']


@pytest.mark.asyncio
async def test_financial_history_pages_are_closed_after_role_change(ui,monkeypatch):
    ui.repo.grant_shop_access(103,ui.shop.id,'accountant')
    rows=[{'action_key':f'finance:test:{i}','category':'finance','title':'PRIVATE '+str(i)+'x'*300,
        'as_of_date':'2026-10-03','priority':2} for i in range(25)]
    monkeypatch.setattr(ui.repo,'action_history',lambda *a,**kw:rows)
    await press(ui,'/action_history',user=103);mid=ui.telegram.sequence
    assert ui.repo.paged_report(ui.bot.id,103,mid)['capability']=='finance'
    ui.repo.grant_shop_access(103,ui.shop.id,'manager');before=len(ui.telegram.methods)
    await callback(ui,mid,'report_page:1',user=103)
    assert [type(m).__name__ for m in ui.telegram.methods[before:]]==['AnswerCallbackQuery']


@pytest.mark.asyncio
async def test_cost_import_rechecks_permission_after_file_download(ui,monkeypatch):
    ui.repo.grant_shop_access(103,ui.shop.id,'accountant')
    await press(ui,'/import_costs',user=103)
    ui.bot.get_file=AsyncMock(return_value=SimpleNamespace(file_path='file'))
    async def download(*a,**kw):ui.repo.grant_shop_access(103,ui.shop.id,'manager')
    ui.bot.download_file=AsyncMock(side_effect=download)
    importer=AsyncMock(side_effect=AssertionError('Revoked cost import'));monkeypatch.setattr('app.bot.handlers.import_costs',importer)
    ui.update_id+=1
    message=types.Message(message_id=ui.update_id,date=ui.telegram.messages[(103,1001)].date,
        chat=types.Chat(id=103,type='private'),from_user=types.User(id=103,is_bot=False,first_name='Employee'),
        document=types.Document(file_id='file',file_unique_id='unique',file_name='costs.csv'))
    await ui.dp.feed_update(ui.bot,types.Update(update_id=ui.update_id,message=message))
    importer.assert_not_called()
    assert 'Недостаточно прав' in last_text(ui,103)


@pytest.mark.asyncio
async def test_old_retry_button_cannot_be_used_by_shop_owner(ui,monkeypatch):
    ui.repo.grant_shop_access(103,ui.shop.id,'owner')
    await press(ui,HOME,user=103);mid=active(ui,103)
    requeue=AsyncMock(side_effect=AssertionError('Not system owner'));monkeypatch.setattr(ui.repo,'requeue_retry_job',requeue)
    await callback(ui,mid,'retry:run:8',user=103)
    requeue.assert_not_called()


@pytest.mark.asyncio
async def test_export_handler_uses_current_role_for_both_formats(ui,monkeypatch,tmp_path):
    captures=[]
    def export(*args,**kwargs):
        captures.append(kwargs);args[-1].write_text('test')
        return SimpleNamespace(path=args[-1],start='2026-09-28',end='2026-09-28')
    monkeypatch.setattr('app.bot.handlers.export_xlsx',export)
    monkeypatch.setattr('app.bot.handlers.export_csv_zip',export)
    for role,fmt in (('manager','xlsx'),('accountant','csv'),('owner','xlsx')):
        ui.repo.grant_shop_access(103,ui.shop.id,role)
        await press(ui,f'/export 1 {fmt}',user=103)
        assert captures[-1]=={'include_finance':role!='manager','include_technical':False}
        assert isinstance(ui.telegram.methods[-1],SendDocument)


@pytest.mark.asyncio
async def test_old_unclassified_snapshots_must_be_reopened_by_employees(ui):
    ui.repo.grant_shop_access(103,ui.shop.id,'owner')
    await press(ui,'/my_access',user=103);message=ui.telegram.messages[(103,ui.update_id)].as_(ui.bot)
    sent=await PagedReportController(ui.ctx).show(message,'Old report\n'+'b'*5000)
    with ui.repo.db.connect() as c:c.execute("UPDATE telegram_paged_reports SET capability='legacy'")
    before=len(ui.telegram.methods);await callback(ui,sent.message_id,'report_page:1',user=103)
    assert [type(m).__name__ for m in ui.telegram.methods[before:]]==['AnswerCallbackQuery']


def action_center(shop_id):
    return ActionCenter(shop_id,date(2026,10,3),(
        ActionItem('finance:missing-cost',2,'finance','PRIVATE FINANCE','Hidden cost','',''),
        ActionItem('system:dead-retries',1,'system','PRIVATE TECH','Hidden path','',''),
        ActionItem('stock:low:123',1,'stock','PUBLIC STOCK','Need stock','',''),
        ActionItem('alert:api_stale:ozon',2,'system','PUBLIC FRESHNESS','Refresh data','','')))


@pytest.mark.asyncio
async def test_actions_old_callbacks_history_and_ack_filter_role(ui,monkeypatch):
    ui.repo.grant_shop_access(103,ui.shop.id,'manager');center=action_center(ui.shop.id)
    for module in ('app.bot.handlers','app.bot.operational_cards'):
        monkeypatch.setattr(module+'.build_action_center',lambda *a,**kw:center)
    rows=[{'action_key':i.action_key,'category':i.category,'title':i.title,'as_of_date':'2026-10-03',
        'priority':i.priority,'current_status':'open'} for i in center.items]
    monkeypatch.setattr(ui.repo,'action_history',lambda *a,**kw:rows)
    monkeypatch.setattr(ui.repo,'set_action_status',AsyncMock(side_effect=AssertionError('closed mutation')))
    await press(ui,'/actions',user=103);mid=ui.telegram.sequence
    text=ui.telegram.messages[(103,mid)].text
    assert 'PUBLIC STOCK' in text and 'PUBLIC FRESHNESS' in text
    assert 'PRIVATE' not in text
    await callback(ui,mid,'action:list',user=103)
    assert 'PRIVATE' not in ui.telegram.messages[(103,mid)].text
    for action in ('view','ack','snooze'):
        before=len(ui.telegram.methods)
        await callback(ui,mid,f'action:{action}:{action_ref("finance:missing-cost")}',user=103)
        assert [type(m).__name__ for m in ui.telegram.methods[before:]]==['AnswerCallbackQuery']
    await callback(ui,mid,f'ops:actions:view:{ui.shop.id}:0:{action_ref("finance:missing-cost")}',user=103)
    assert 'PRIVATE' not in ui.telegram.messages[(103,mid)].text
    await press(ui,'/action_history',user=103)
    assert 'PRIVATE' not in last_text(ui,103)
    await press(ui,'/action_ack finance:missing-cost',user=103)
    await press(ui,'/action_snooze system:dead-retries',user=103)
    ui.repo.set_action_status.assert_not_called()


@pytest.mark.asyncio
async def test_forged_help_technical_section_and_finance_picker_are_blocked(ui):
    ui.repo.grant_shop_access(103,ui.shop.id,'manager')
    await press(ui,HOME,user=103);mid=active(ui,103)
    before=len(ui.telegram.methods)
    await callback(ui,mid,f'mh:{ui.shop.id}:103:technical:section:show',user=103)
    assert [type(m).__name__ for m in ui.telegram.methods[before:]]==['AnswerCallbackQuery']
    ui.repo.grant_shop_access(103,ui.shop.id,'accountant')
    await press(ui,COMMAND_BUTTONS['finance'],user=103);mid=active(ui,103)
    button=inline_data(ui,mid,'Вчера',user=103)
    ui.repo.grant_shop_access(103,ui.shop.id,'manager')
    before=len(ui.telegram.methods);await callback(ui,mid,button,user=103)
    assert [type(m).__name__ for m in ui.telegram.methods[before:]]==['AnswerCallbackQuery']


@pytest.mark.asyncio
async def test_person_without_access_gets_own_copyable_id(ui):
    async def context_for_user(uid):raise PermissionError
    registry=SimpleNamespace(maintenance_lock=asyncio.Lock(),context_for_user=context_for_user)
    ui.dp=Dispatcher();ui.dp.message.outer_middleware(ShopContextMiddleware(registry));register_handlers(ui.dp,ui.ctx)
    await press(ui,'/start',user=777)
    assert '<code>777</code>' in last_text(ui,777)
    assert 'Shop' not in last_text(ui,777)


def seed_export(ui):
    conn=ui.repo.ensure_connection(ui.shop.id,'ozon','Ozon')
    product=ui.repo.ensure_product(ui.shop.id,'P','PUBLIC PRODUCT',cost_price=876543.21)
    listing=ui.repo.ensure_listing(product.id,conn.id,'123')
    run=ui.repo.record_success(conn.id,'combined','2026-09-28',{},[
        MetricPoint(conn.id,'2026-09-28','ordered_units',5,'units'),
        MetricPoint(conn.id,'2026-09-28','ordered_revenue',100,'RUB'),
        MetricPoint(conn.id,'2026-09-28','ad_spend',10,'RUB'),
        MetricPoint(conn.id,'2026-09-28','marketplace_net',987654.32,'RUB'),
        MetricPoint(conn.id,'2026-09-28','commission',999888.77,'RUB')])
    ui.repo.save_product_metrics(run,[ProductMetricPoint(listing.id,'2026-09-28','ordered_units',5,'units')])


@pytest.mark.parametrize('fmt',['xlsx','csv'])
def test_manager_export_no_finances_hidden_fields_or_technical_tasks(ui,tmp_path,monkeypatch,fmt):
    seed_export(ui)
    center=action_center(ui.shop.id)
    monkeypatch.setattr('app.services.exporting.build_action_center',lambda *a,**kw:center)
    original=ui.repo.product_period_totals
    def product_rows(*a,**kw):return [dict(row,cost_price=876543.21) for row in original(*a,**kw)]
    monkeypatch.setattr(ui.repo,'product_period_totals',product_rows)
    exporter=export_xlsx if fmt=='xlsx' else export_csv_zip
    result=exporter(ui.repo,ui.shop.id,date(2026,9,28),1,tmp_path/('out.xlsx' if fmt=='xlsx' else 'out.zip'),
        include_finance=False,include_technical=False)
    if fmt=='xlsx':
        wb=load_workbook(result.path);names=wb.sheetnames
        content=json.dumps([[list(row) for row in sheet.values] for sheet in wb],ensure_ascii=False,default=str)
    else:
        with zipfile.ZipFile(result.path) as archive:
            names=[p.removesuffix('.csv').title() for p in archive.namelist()]
            content='\n'.join(archive.read(p).decode('utf-8-sig') for p in archive.namelist())
    assert not {'Finance','Costs','Management','Reconciliation',
                'Финансовая сводка','Себестоимость','Результат магазина','Сверка данных'}&set(names)
    assert 'PUBLIC PRODUCT' in content and 'PUBLIC STOCK' in content and 'ad_spend' in content
    assert all(hidden not in content for hidden in ('987654.32','876543.21','999888.77','cost_price','PRIVATE','credential_profile','schema_version'))


def test_accountant_export_keeps_finances_but_removes_technical(ui):
    seed_export(ui)
    tables=collect_export_tables(ui.repo,ui.shop.id,date(2026,9,28),1,include_finance=True,include_technical=False)
    assert {'Finance','Costs','Management','Reconciliation'}<=set(tables)
    assert any(row['value']==987654.32 for row in tables['Finance'])
    assert not any(row['key']=='credential_profile' for row in tables['Summary'])


def test_schema19_role_upgrade_preserves_access_timestamps_and_snapshots(tmp_path):
    from app.storage.database import MIGRATIONS
    db=Database(tmp_path/'old.sqlite3')
    with db.connect() as c:
        c.execute('CREATE TABLE schema_migrations (version INTEGER PRIMARY KEY,applied_at TEXT NOT NULL)')
        for version,migration in MIGRATIONS.items():
            if version==20:break
            migration(c);c.execute("INSERT INTO schema_migrations VALUES(?,datetime('now'))",(version,))
    repo=Repository(db);shop=repo.ensure_shop(repo.ensure_seller(1).id)
    for uid,role in ((101,'owner'),(102,'viewer'),(103,'owner')):repo.grant_shop_access(uid,shop.id,role)
    with db.connect() as c:
        now='2026-10-03T06:00:00Z'
        c.execute('INSERT INTO bot_users VALUES(?,?,?,?,?)',(104,'Analyst',1,now,now))
        c.execute('INSERT INTO user_shop_access VALUES(?,?,?,?,?)',(104,shop.id,'analyst',now,now))
        previous=[tuple(row) for row in c.execute('SELECT * FROM user_shop_access ORDER BY telegram_user_id')]
    repo.save_report_card(123,101,9,shop_id=shop.id,report_day='2026-10-01',summary_html='Summary',
        details_html='Detail',accruals_html='Accruals',section='summary',status_note='',page=0)
    from app.storage import LATEST_SCHEMA_VERSION
    assert db.initialize_safely()==LATEST_SCHEMA_VERSION and db.integrity_check()
    with db.connect() as c:
        migrated=[tuple(row) for row in c.execute('SELECT * FROM user_shop_access ORDER BY telegram_user_id')]
    assert migrated==[(uid,sid,'accountant' if role=='analyst' else role,created,updated)
        for uid,sid,role,created,updated in previous]
    assert repo.report_card(123,101,9)['summary_html']=='Summary'
    assert repo.can_user(102,shop.id,'finance') and not repo.can_user(102,shop.id,'operate')
    assert repo.can_user(104,shop.id,'costs')
